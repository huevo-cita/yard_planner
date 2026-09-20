#!/usr/bin/env python3
"""Prove the calendar refuses to render over a plan it no longer agrees with.

    python3 tools/test_week.py
    python3 tools/test_week.py -v        show every check's detail

This repo tests its guardrails rather than asserting them — `test_gate.py` for
the doubt gate, `test_changelog.py` for the prose lint — because a mechanism
hand-tested once in a chat that no longer exists is not evidence of anything.

`lib.week` is the third guardrail and it turns entirely on one claim: that a date
changed in `PLAN.md` and not in `tasks.json` is caught rather than silently
disagreeing. The properties worth proving, each easy to pass by accident:

  the digest fires     editing a source section makes the render refuse, naming
                       that section. This is the layer that catches a task
                       nobody ever transcribed, which comparing dates cannot
  source means source  a section named under `source` is hashed whatever file it
                       is in, and one named under `reference` or `technique` is
                       not. Both halves matter: the first is the hole the g05
                       gutter programme took its dates through, and the second
                       is what keeps the check from firing on a correct record
  the date check fires moving a date the record cites makes the render refuse,
                       naming the task. Separate from the digest because a
                       digest says only that something moved
  it is not noisy      a clean record produces no findings at all. A check that
                       fires on a correct document is the one that gets switched
                       off, and then there is no check. Tested explicitly,
                       because the obvious implementation — scanning prose for
                       dates and asking which are missing — fails exactly here
  no false positive    a date the plan writes as a range start, "Wed 2 - Sat 5
                       Sep", is stated and must not read as missing
  the flag is honoured a date the plan never states is fine when the task says
                       so with `date_inferred`, and not otherwise
  force stamps         `--force` renders, and the page names what it came past.
                       "Provisional" on its own is not something anyone can act
                       on, so the stamp has to carry the section and the task
  restamp is not a fix `--restamp` clears the digest but must not clear a date
                       the plan no longer states, or the escape hatch swallows
                       the finding it exists to record

Everything runs against a temporary GARDEN_ROOT, so no real yard is read or
written and none of this touches personal data.
"""

import argparse
import datetime
import json
import os
import shutil
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

SLUG = "testyard"

results = []          # (state, label, detail) — state in pass/FAIL
verbose = False


def record(state, label, detail=""):
    results.append((state, label, detail))
    print(f"  {'ok  ' if state == 'pass' else 'FAIL'}  {label}")
    if detail and (state != "pass" or verbose):
        for line in str(detail).strip().splitlines():
            print(f"          {line}")


def ok(cond, label, detail=""):
    record("pass" if cond else "FAIL", label, "" if cond else detail)


# --------------------------------------------------------------- the fixture

# Small, but shaped like the real thing: two numbered sections, a date written
# plainly, a date written as the start of a dashed range with the month only at
# the end, and a date the prose never states at all. Each exists to make one
# check above fail if it is wrong.
PLAN = """# Scratch — the plan

## 1. The next seven days

- **Mon 31 Aug, 10 min — phone the nursery.**
- **Wed 2 - Sat 5 Sep — watch the tray.**

## 2. Weekend by weekend

### Sep 6-13

- **Sat 12 Sep, 2 h — prune the roses.**
- Weekday evening, 30 min — set out the toad abodes.
"""

# Not a plan document, and numbered the way a long reference document is: a
# parent section with dotted subsections and no `.` after the number. The whole
# g05 gutter programme came out of a file shaped exactly like this, cited it
# under `source`, and got no digest because the filename was not on a whitelist.
RESEARCH = """# Scratch — research

## 4. Slope, and why it matters

Body of section four.

### 4.1 Read this before buying anything

Body of section four point one.

### 4.2 The other one

Body of section four point two.
"""

TASKS = {
    "yard": SLUG,
    "schema_version": 1,
    "target_date": "2026-12-13",
    "sources": {"PLAN.md#1": {"digest": None, "read": None},
                "PLAN.md#2": {"digest": None, "read": None}},
    "suppliers": {},
    "shopping": [],
    "tasks": [
        {"id": "t001", "date": "2026-08-31", "minutes": 10, "kind": "call",
         "title": "Phone the nursery", "where": {"place": "phone"},
         "source": ["PLAN.md#1"], "done": False},
        {"id": "t002", "window": ["2026-09-02", "2026-09-05"], "minutes": 1,
         "kind": "inspect", "title": "Watch the tray",
         "where": {"place": "indoors"}, "source": ["PLAN.md#1"], "done": False},
        {"id": "t003", "date": "2026-09-12", "minutes": 120, "kind": "prune",
         "title": "Prune the roses", "where": {"bed": "roses"},
         "source": ["PLAN.md#2"], "done": False},
        {"id": "t004", "date": "2026-09-11", "minutes": 30, "kind": "build",
         "title": "Set out the toad abodes", "where": {"place": "the outfalls"},
         "date_inferred": True, "date_note": "PLAN.md says 'weekday evening'",
         "source": ["PLAN.md#2"], "done": False},
    ],
}


def make_yard(root):
    d = os.path.join(root, SLUG)
    os.makedirs(d, exist_ok=True)
    with open(os.path.join(d, "PLAN.md"), "w") as f:
        f.write(PLAN)
    with open(os.path.join(d, "research-scratch.md"), "w") as f:
        f.write(RESEARCH)
    with open(os.path.join(d, "tasks.json"), "w") as f:
        json.dump(TASKS, f, indent=2)
    return d


def edit_plan(yard, old, new):
    p = os.path.join(yard, "PLAN.md")
    with open(p) as f:
        text = f.read()
    assert old in text, f"fixture does not contain {old!r}"
    with open(p, "w") as f:
        f.write(text.replace(old, new, 1))


def kinds(findings):
    return sorted({f["kind"] for f in findings})


def subjects(findings, kind):
    return sorted(f["subject"] for f in findings if f["kind"] == kind)


# ------------------------------------------------------------------ the checks

def check_clean(week, yard):
    """A correct record produces nothing. The check that keeps the rest usable."""
    week.restamp(SLUG)
    found = week.check(SLUG)
    ok(not found, "a record that agrees with the plan reports nothing",
       "\n".join(f["message"] for f in found))

    # t002 is written "Wed 2 - Sat 5 Sep" and the month appears once, at the end.
    # A naive matcher looks for "2 Sep", does not find it, and reports a date the
    # plan states plainly as missing.
    ok(not any(f["subject"] == "t002" for f in found),
       "a date written as the start of a dashed range is not reported missing")

    # t004's date is nowhere in the prose, and the record says so.
    ok(not any(f["subject"] == "t004" for f in found),
       "a date flagged date_inferred is not reported missing")


def check_unstamped(week, yard):
    """Never stamped is not the same as agreeing, and must not read as clean."""
    data = week.load(SLUG)
    data["sources"]["PLAN.md#1"]["digest"] = None
    week.save(SLUG, data)
    found = week.check(SLUG)
    ok("unstamped" in kinds(found) and subjects(found, "unstamped") == ["PLAN.md#1"],
       "a section that was never stamped is reported, not assumed fine",
       [f["message"] for f in found])
    week.restamp(SLUG)


def check_uncited(week, yard):
    """A section a task leans on but nothing hashes is a hole in the digest."""
    data = week.load(SLUG)
    data["tasks"][0]["source"].append("PLAN.md#9")
    week.save(SLUG, data)
    found = week.check(SLUG)
    ok(subjects(found, "uncited") == ["PLAN.md#9"],
       "a section cited by a task but carrying no digest is reported",
       [f["message"] for f in found])
    data = week.load(SLUG)
    data["tasks"][0]["source"] = ["PLAN.md#1"]
    week.save(SLUG, data)


def check_source_is_hashed_whatever_the_file(week, yard):
    """`source` means extracted from, so it is hashed even in a research file.

    This is the hole the g05 gutter walked through. Seven tasks took their
    content from `research-guttering.md`, named it under `source`, and got no
    digest, because the check only asked about a fixed list of plan documents.
    Section 5 of that file then turned out to rest on a driveway at the wrong
    corner of the house and nothing anywhere noticed.
    """
    data = week.load(SLUG)
    data["tasks"][0]["source"].append("research-scratch.md#4.1")
    week.save(SLUG, data)
    found = week.check(SLUG)
    ok(subjects(found, "uncited") == ["research-scratch.md#4.1"],
       "a research section cited under `source` is reported as unhashed",
       [f["message"] for f in found])

    # And the escape hatch works in one command rather than a hand edit.
    week.restamp(SLUG)
    ok("research-scratch.md#4.1" in week.load(SLUG).get("sources", {}),
       "--restamp adopts a cited section that carried no digest")
    ok(not week.check(SLUG), "and the record is clean once it has",
       [f["message"] for f in week.check(SLUG)])

    # The point of hashing it: a change to it now surfaces.
    p = os.path.join(yard, "research-scratch.md")
    with open(p) as f:
        before = f.read()
    with open(p, "w") as f:
        f.write(before.replace("Body of section four point one.",
                               "Actually the opposite, and it was never true."))
    found = week.check(SLUG)
    ok("research-scratch.md#4.1" in subjects(found, "stale"),
       "editing that research section makes the digest stale, naming it",
       [f["message"] for f in found])
    with open(p, "w") as f:
        f.write(before)

    data = week.load(SLUG)
    data["tasks"][0]["source"] = ["PLAN.md#1"]
    week.save(SLUG, data)


def check_reference_is_not_hashed(week, yard):
    """`reference` and `technique` are links, not provenance, and stay unhashed.

    The other half of the rule, and the one that keeps it usable. If every
    document a task points at had to be stamped, a task linking three technique
    notes would make the check fire on a correct record, which is how a check
    gets switched off.
    """
    data = week.load(SLUG)
    data["sources"].pop("research-scratch.md#4.1", None)
    data["tasks"][0]["reference"] = "research-scratch.md#4.2"
    data["tasks"][0]["technique"] = "research-scratch.md#4"
    week.save(SLUG, data)
    found = week.check(SLUG)
    ok(not found, "a research file under `reference` or `technique` is not hashed",
       [f["message"] for f in found])
    data = week.load(SLUG)
    data["tasks"][0].pop("reference"), data["tasks"][0].pop("technique")
    week.save(SLUG, data)


def check_dotted_anchor(week, yard):
    """`#4.1` finds `### 4.1 ...`, and `#4` does not swallow its own children.

    A reference document numbers subsections `4.1` and writes no `.` after the
    number. Before this the anchor fell through to a slug match, `4.1` was not a
    substring of `41-read-this-before-buying-anything`, and the reference
    resolved to nothing — so a task could cite a section that could never be
    hashed and the failure looked like a typo.
    """
    sec, err = week.resolve(yard, "research-scratch.md#4.1")
    ok(sec is not None and sec["text"].startswith("4.1 Read"),
       "a dotted section anchor resolves", err or (sec or {}).get("text"))

    sec, err = week.resolve(yard, "research-scratch.md#4")
    ok(sec is not None and sec["text"].startswith("4. Slope"),
       "the parent anchor still resolves to the parent, not to a child",
       err or (sec or {}).get("text"))

    sec, err = week.resolve(yard, "research-scratch.md#4.9")
    ok(sec is None, "a dotted anchor that names no heading is an error, not a guess",
       (sec or {}).get("text"))


def check_digest(week, yard):
    """Editing a section makes the render refuse, naming that section.

    The edit deliberately adds a task rather than moving one, because that is
    the case no date comparison can catch: nothing in tasks.json is wrong, and
    something in the plan was never transcribed.
    """
    edit_plan(yard, "### Sep 6-13",
              "### Sep 6-13\n\n- **Sun 13 Sep, 1 h — mulch the roses.**")
    found = week.check(SLUG)
    ok(subjects(found, "stale") == ["PLAN.md#2"],
       "a task added to the plan and never transcribed makes the digest stale",
       [f["message"] for f in found])

    try:
        week.calendar(SLUG)
        rendered = True
    except SystemExit as exc:
        rendered = False
        message = str(exc)
    ok(not rendered, "--calendar refuses while a section is stale")
    if not rendered:
        ok("PLAN.md §2" in message,
           "the refusal names the section that moved", message)
        ok("--force" in message and "--check" in message,
           "the refusal says what to do next", message)


def check_date(week, yard):
    """Moving a date the record cites is caught, and names the task."""
    week.restamp(SLUG)                      # the digest is no longer the finding
    edit_plan(yard, "Sat 12 Sep", "Sat 19 Sep")
    week.restamp(SLUG)                      # and neither is the second edit
    found = week.check(SLUG)
    ok(subjects(found, "date") == ["t003"],
       "a date moved in the plan is reported against the task that holds it",
       [f["message"] for f in found])
    ok(any("12 September" in f["message"] for f in found if f["kind"] == "date"),
       "the finding says which date the record still believes",
       [f["message"] for f in found])

    # This is the property that makes --restamp safe to offer: it settles the
    # digest, which is a claim about having re-read the section, and settles
    # nothing about a date that is genuinely wrong.
    week.restamp(SLUG)
    still = week.check(SLUG)
    ok(subjects(still, "date") == ["t003"],
       "--restamp does not clear a date the plan no longer states",
       [f["message"] for f in still])


def check_force(week, yard):
    """It renders, and the page names what it came past."""
    found = week.check(SLUG)
    try:
        path, _, text = week.calendar(SLUG, force=True)
    except SystemExit as exc:
        record("FAIL", "--force renders anyway", str(exc))
        return
    record("pass", "--force renders anyway")

    head = text.split("\n## ", 1)[0]
    ok("PROVISIONAL" in head, "the page is stamped provisional", head[:400])
    ok("t003" in head,
       "the stamp names the task whose date the plan no longer states", head[:400])
    ok(len(found) == 1 or "PLAN.md §" in head,
       "the stamp names the section, where a section moved", head[:400])
    ok(head.count("PROVISIONAL") == 1,
       "the stamp appears once, at the top, not per week")

    # A stamp that only said "provisional" would pass every check above except
    # this one: the point is that a reader can tell which dates to distrust.
    ok(week.stamp(found) and week.stamp(found) != "provisional",
       "the stamp is specific rather than a bare word", week.stamp(found))


def check_cheap_paths(week, yard):
    """The ways out have to stay open, or the refusal is a dead end."""
    for label, fn in (("--check", lambda: week.check(SLUG)),
                      ("the terminal week view", lambda: week.report(SLUG)),
                      ("--shop", lambda: week.shop(SLUG, 3))):
        try:
            fn()
            record("pass", f"{label} still runs while the record is refusing")
        except SystemExit as exc:
            record("FAIL", f"{label} still runs while the record is refusing",
                   str(exc))


def check_sync(week, yard):
    """Ticks come back, through the escaping the Docs export applies."""
    export = os.path.join(yard, "exported.md")
    with open(export, "w") as f:
        f.write("> - [x] **Phone the nursery** · 10 min · phone\n"
                "> - [ ] **Watch the tray** · 1 min · indoors\n"
                "> - [x] **Prune the roses** · 2 h · roses\n")
    changed, unmatched, seen = week.sync(SLUG, export)
    ok(seen == 3, f"every checkbox in the export is read (saw {seen})")
    ok(sorted(c[0] for c in changed) == ["t001", "t003"],
       "the ticked ones, and only those, are marked done", changed)
    ok(not unmatched, "every item in the export matches a task", unmatched)

    # The export backslash-escapes punctuation, so a title with a dash in it is
    # the case that silently matches nothing.
    data = week.load(SLUG)
    data["tasks"][3]["title"] = "Set out the toad abodes - four of them"
    week.save(SLUG, data)
    with open(export, "w") as f:
        f.write("> - [x] **Set out the toad abodes \\- four of them** · 30 min\n")
    changed, unmatched, _ = week.sync(SLUG, export)
    ok(sorted(c[0] for c in changed) == ["t004"] and not unmatched,
       "a title the export escaped still matches its task", (changed, unmatched))
    check_done_prefix(week, yard)
    check_struck_out(week, yard)


def check_struck_out(week, yard):
    """Striking a line out in the Doc is how a person marks it done by hand.

    It exports as `~~` around the whole item, and it lands in two places that
    fail differently. On an ordinary ticked line the tildes ride along on the
    front of the title, so the task is reported as drift and its tick is lost —
    silently, because drift is reported as a formatting problem rather than as a
    completed job. On a line the publish step already marked done, the tildes sit
    in front of `DONE ·` so the marker never matches, and every such task keys on
    the word DONE and collides with the others.
    """
    data = week.load(SLUG)
    for t in data["tasks"][:3]:
        t["done"] = False
    week.save(SLUG, data)
    titles = [t["title"] for t in data["tasks"][:3]]

    export = os.path.join(yard, "struck-export.md")
    with open(export, "w") as f:
        f.write(f"> - [x] ~~**{titles[0]}** · 10 min · phone~~\n"
                f"> - [x] ~~DONE · **{titles[1]}** · 1 min · indoors~~\n"
                f"> - [ ] **{titles[2]}** · 2 h · roses\n")
    changed, unmatched, seen = week.sync(SLUG, export)
    ok(seen == 3, f"a struck-out line is still one checkbox item, keyed on its "
                  f"own title (saw {seen} of 3)")
    ok(not unmatched, "a struck-out item is not reported as drift", unmatched)
    ok(sorted(c[0] for c in changed) == ["t001", "t002"]
       and all(c[2] is True for c in changed),
       "both a struck-out tick and a struck-out DONE marker come back as done",
       changed)


def check_done_prefix(week, yard):
    """A done task publishes as "- [ ] DONE · title" and has to survive the trip.

    The .docx import cannot carry a ticked box, so `publish` encodes doneness in
    the text and sends the box out empty. Both halves of reading that back can
    fail, and they fail in opposite directions: take the title without stripping
    the marker and every done task keys on the word DONE, so they collide and
    the count comes back short; strip it but read the empty box at face value
    and every finished task is marked undone, which erases the record.
    """
    data = week.load(SLUG)
    for t, done in zip(data["tasks"], (True, True, False)):
        t["done"] = done
    week.save(SLUG, data)
    titles = [t["title"] for t in data["tasks"][:3]]

    export = os.path.join(yard, "done-export.md")
    with open(export, "w") as f:
        f.write(f"> - [ ] DONE · **{titles[0]}** · 10 min · phone\n"
                f"> - [ ] DONE · **{titles[1]}** · 1 min · indoors\n"
                f"> - [ ] **{titles[2]}** · 2 h · roses\n")
    changed, unmatched, seen = week.sync(SLUG, export)
    ok(seen == 3, f"two DONE lines key separately rather than colliding "
                  f"(saw {seen} of 3)")
    ok(not unmatched, "no phantom 'DONE' item is reported as drift", unmatched)
    ok(not changed, "a task already done stays done through the round trip",
       changed)

    # And the other direction: the marker is what says done, not the box.
    data = week.load(SLUG)
    for t in data["tasks"][:2]:
        t["done"] = False
    week.save(SLUG, data)
    changed, _, _ = week.sync(SLUG, export)
    ok(sorted(c[0] for c in changed) == ["t001", "t002"]
       and all(c[2] is True for c in changed),
       "the DONE marker marks a task done even though its box is empty",
       changed)


def check_column_widths():
    """A table the importer autofits is mostly scroll, so the widths are pinned.

    The shopping table is the case this is for: one prose column beside four
    holding a date, a price and a name. Shared out evenly the prose wraps to
    dozens of lines and every other cell in the row is white space.
    """
    from lib import builddoc

    def cells(*texts):
        return [[(t, False, False)] for t in texts]

    # One prose column against three short ones, the shopping table's shape.
    rows = [cells("Buy", "By", "Cost", "Ask for"),
            cells("Fire ant bait", "Sun 6 Sep", "$22-30", "x" * 2000),
            cells("Dill seed", "Sat 5 Sep", "$3-6", "y" * 400)]
    w = builddoc.column_widths(rows, 4)

    ok(abs(sum(w) - builddoc.TABLE_WIDTH_IN) < 0.01,
       "the columns add up to the text column exactly", sum(w))
    ok(min(w) >= builddoc.MIN_COL_IN - 1e-9,
       "no column falls under the floor a date needs to stay on one line", w)
    ok(w[3] == max(w) and w[3] > builddoc.TABLE_WIDTH_IN / 4,
       "the prose column takes more than an even share", w)
    # The damping is the point: raw proportion would give the prose column
    # 2000/2050 of the table and leave the date under a tenth of an inch.
    ok(w[3] < builddoc.TABLE_WIDTH_IN * 0.75,
       "and not so much that the short columns are unreadable", w)
    ok(w[1] > w[2],
       "a longer short column still beats a shorter one", (w[1], w[2]))

    # Degenerate: more columns than the page can hold falls back to even.
    many = builddoc.column_widths([cells(*"abcdefghijklmnopqrst")], 20)
    ok(abs(sum(many) - builddoc.TABLE_WIDTH_IN) < 0.01
       and len(set(round(x, 6) for x in many)) == 1,
       "more columns than the floor allows shares the width evenly instead of "
       "overflowing the page", (len(many), sum(many)))

    # The widths have to reach w:tblGrid, not just w:tcW. Google Docs reads the
    # grid and ignores the cells, so a version that set only the cells passed
    # every check above and still imported as five even columns.
    from docx.oxml.ns import qn
    md = os.path.join(tempfile.mkdtemp(prefix="yard-doc-test-"), "t.md")
    with open(md, "w") as fh:
        fh.write("| Buy | By | Ask for |\n|---|---|---|\n"
                 "| bait | Sun 6 Sep | %s |\n" % ("x" * 900))
    out = md[:-3] + ".docx"
    imgs = {}
    html = builddoc.build_html(md, imgs)
    from docx import Document
    doc = Document()
    builddoc.Conv(doc, imgs).feed(html)
    doc.save(out)

    t = Document(out).tables[0]
    gridEl = t._tbl.find(qn("w:tblGrid"))
    grid = [int(c.get(qn("w:w"))) / 1440.0
            for c in gridEl.findall(qn("w:gridCol"))] if gridEl is not None else []
    ok(len(grid) == 3 and len(set(round(g, 3) for g in grid)) == 3,
       "w:tblGrid carries three different widths, not add_table's even split", grid)
    ok(grid and grid[2] == max(grid),
       "and the widest grid column is the prose one", grid)
    layout = t._tbl.tblPr.find(qn("w:tblLayout"))
    ok(layout is not None and layout.get(qn("w:type")) == "fixed",
       "the table layout is fixed, so nothing re-fits it on import",
       layout if layout is None else layout.get(qn("w:type")))
    shutil.rmtree(os.path.dirname(md), ignore_errors=True)


# ------------------------------------------------- the one navigable set

# A second yard, because these checks want a design, a call and two shapes of
# placement, and bending the first fixture to carry them would weaken the
# checks it already makes.
SCREENS = "testyard-screens"

S_PLAN = """# Screens — the plan

## 1. Sowing

How to sow.

## 2. Pests

What eats what.

### Sources

One list of sources.

## Sources

Another heading with the same words, lower down.
"""

S_DESIGN = {
    "yard": SCREENS,
    "plants": [
        {"zone": "bed_a", "name": "Cedar sedge", "count": 5,
         "botanical": "Carex planostachys", "note": "dry shade"},
        {"zone": "bed_a", "name": "Gulf muhly", "count": 3,
         "botanical": "Muhlenbergia capillaris"},
        {"zone": "bed_a", "name": "Pale-leaf yucca", "count": 1,
         "botanical": "Yucca pallida"},
        {"zone": "bed_a", "name": "Milkweed - ASK FOR Asclepias tuberosa",
         "count": 2, "botanical": "Asclepias tuberosa / A. asperula"},
        {"zone": "bed_a", "name": "Rosemary 'Tuscan Blue' (upright)",
         "count": 1, "botanical": "Salvia rosmarinus"},
        {"zone": "bed_raised", "name": "Lettuce (mesclun)", "count": 3},
    ],
    "cut_and_why": [{"item": "Autumn sage x3", "reason": "area, not light"}],
}

S_SOURCING = {"yard": SCREENS, "suppliers": [
    {"id": "nursery", "name": "A Nursery", "address": "1 Road",
     "phone": "(512) 555-0100", "hours": "9-5", "distance_mi": 4.1}]}

S_TASKS = {
    "yard": SCREENS, "schema_version": 1, "sources": {}, "suppliers": {},
    "shopping": [{"id": "b01", "item": "The plant order", "supplier": "nursery",
                  "by": "2026-09-25", "cost_usd": [100, 120],
                  "confidence": "estimated", "source": [],
                  "ask": ("Yucca pallida NOT Y. recurvifolia, which trunks. "
                          "Asclepias tuberosa or A. asperula by name - "
                          "anything sold as milkweed here is A. curassavica. "
                          "No Autumn sage, which left the design.")}],
    "tasks": [
        {"id": "t101", "date": "2026-09-21", "minutes": 20, "kind": "call",
         "title": "Phone the nursery", "where": {"place": "phone",
                                                 "supplier": "nursery"},
         "buy": ["b01"], "source": [], "done": False,
         "how": ["Ask for Carex planostachys and DO NOT ACCEPT 'Texas sedge', "
                 "which is C. texensis, the wrong plant for dry shade",
                 "Same trap as the Yucca above: the label is not the plant"]},
        {"id": "t102", "date": "2026-10-24", "minutes": 120, "kind": "plant",
         "title": "Plant bed a", "source": [], "done": False,
         "technique": "PLAN.md#1", "reference": "PLAN.md section 2",
         "doubts": ["d01", "d02"],
         "where": {"bed": "bed a", "placements": [
             {"bed": "a", "at": "ft 1", "plant": "Gulf muhly, cedar sedge x2"},
             {"bed": "a", "at": "ft 6", "plant": "Texas sedge x2"},
             {"bed": "a", "at": "ft 9", "plant": "Rosemary"}]}},
        {"id": "t103", "date": "2026-11-01", "minutes": 20, "kind": "harvest",
         "title": "Harvest the bed", "source": [], "done": False,
         "where": {"bed": "raised bed", "placements": [
             {"squares": ["1-1"], "plant": "Cut the mesclun"}]}},
        {"id": "t104", "minutes": 10, "kind": "water", "source": [],
         "title": "Water the new plants", "done": False,
         "where": {"bed": "bed a"},
         "repeat": {"from": "2026-10-19", "to": "2026-11-15", "every": "day"}},
    ],
}

S_DOUBTS = {"yard": SCREENS, "cards": [
    {"id": "d01", "question": "Which Carex does the nursery actually stock?",
     "status": "open", "blocks": ["design"], "effort": "one phone call"},
    {"id": "d02", "question": "Is bed a in shade after noon?",
     "status": "settled", "answer": "yes, from two o'clock"},
]}


def make_screens(root):
    d = os.path.join(root, SCREENS)
    os.makedirs(d, exist_ok=True)
    with open(os.path.join(d, "PLAN.md"), "w") as f:
        f.write(S_PLAN)
    for name, obj in (("design.json", S_DESIGN), ("sourcing.json", S_SOURCING),
                      ("tasks.json", S_TASKS), ("doubts.json", S_DOUBTS)):
        with open(os.path.join(d, name), "w") as f:
            json.dump(obj, f, indent=2)
    return d


def check_one_slug(yard):
    """The anchor a link predicts is the anchor the publisher actually writes.

    This is the property the whole navigable set rests on, and it used to hold
    by accident: two separate slug functions that happened to agree. Proved
    here by publishing the document and looking for the predicted id in it,
    rather than by comparing the two functions to each other — which would
    pass just as happily if both were wrong in the same way.
    """
    from lib import buildhtml, links

    out, _, _ = buildhtml.convert(os.path.join(yard, "PLAN.md"))
    with open(out, encoding="utf-8") as fh:
        page = fh.read()

    for ref in ("PLAN.md#1", "PLAN.md#2"):
        url, sec, err = links.target(yard, ref)
        anchor = (url or "#").split("#", 1)[1]
        ok(not err and f'id="{anchor}"' in page,
           f"{ref} predicts an anchor the published page really has", err or url)

    # Two headings read "Sources". The publisher numbers the second, and a
    # prediction that does not number it lands at the wrong one.
    secs = [s for s in links.anchored(os.path.join(yard, "PLAN.md"))
            if s["text"] == "Sources"]
    ok(len(secs) == 2 and secs[0]["anchor"] == "sources"
       and secs[1]["anchor"] == "sources-1",
       "a repeated heading is numbered the way the publisher numbers it",
       [s["anchor"] for s in secs])
    ok(all(f'id="{s["anchor"]}"' in page for s in secs),
       "and both of those ids are in the page")

    ok(links.target(yard, "PLAN.md#nope")[2],
       "an anchor that matches no heading is an error, not a guess")
    ok(links.target(yard, "https://example.org/x")[0]
       == "https://example.org/x",
       "an external URL passes through rather than being reported broken")


def check_matcher(yard):
    """Every plant a line names, not the first one that happens to match."""
    from lib import plants, yards

    design = yards.load(SCREENS, "design.json")
    idx = plants.index(design)

    hits = [n for n, _ in plants.match("Gulf muhly, cedar sedge x2", idx)]
    ok(hits == ["gulf muhly", "cedar sedge"],
       "a line naming two plants resolves to both, in the order written", hits)

    ok(not plants.match("Texas sedge x2", idx),
       "a name the design does not hold matches nothing, rather than 'sedge'")

    ok([n for n, _ in plants.match("Rosemary", idx)] == ["rosemary"],
       "a cultivar is found by the plain name a person says out loud")

    milkweed = next(p for p in design["plants"] if "Milkweed" in p["name"])
    ok(plants.binomials(milkweed)
       == ["Asclepias tuberosa", "Asclepias asperula"],
       "an abbreviated second species is expanded against its own genus",
       plants.binomials(milkweed))

    ok("milkweed" in idx,
       "a name is indexed with the reader's aside taken off", sorted(idx))


def check_bundle(yard):
    """The join every page renders from, including what it could not join."""
    from lib import bundle

    data = bundle.build(SCREENS)
    t102 = next(t for t in data["tasks"] if t["id"] == "t102")

    good = [l for l in t102["links"] if l["kind"] == "technique"]
    ok(good and good[0]["url"] == "PLAN.html#1-sowing",
       "a resolvable technique becomes a real page anchor",
       good and good[0]["url"])

    bad = [l for l in t102["links"] if l["kind"] == "reference"]
    ok(bad and bad[0]["url"] is None and bad[0]["error"],
       "a reference written as prose is kept, carrying why it failed",
       bad and bad[0]["error"])

    rows = {p["at"]: p for p in t102["placements"]}
    ok(len(rows["ft 1"]["named"]) == 2 and not rows["ft 1"]["unmatched"],
       "a placement carries every plant it names")
    ok(rows["ft 6"]["unmatched"] and rows["ft 6"]["positional"],
       "a bed-and-foot-mark line naming nothing is both unmatched and positional")

    t103 = next(t for t in data["tasks"] if t["id"] == "t103")
    ok(t103["placements"][0]["unmatched"]
       and not t103["placements"][0]["positional"],
       "a square line is unmatched but not positional, so it is not a fault")

    ok(any(d["file"] == "PLAN.md" and d["sections"] for d in data["documents"]),
       "the bundle carries each document's anchors for anything that links in")


def check_link_report(yard):
    """The report names the real disagreement and stays quiet about the rest."""
    from lib import week as W

    found = W.link_check(SCREENS)
    refs = [f for f in found if f["kind"] == "reference"]
    beds = [f for f in found if f["kind"] == "placement"]

    ok(len(refs) == 1 and refs[0]["subject"] == "t102",
       "the reference written as prose is reported once",
       [f["message"] for f in refs])
    ok(len(beds) == 1 and "Texas sedge" in beds[0]["message"],
       "the planting position naming no design plant is reported",
       [f["message"] for f in beds])
    ok(not any(f["subject"] == "t103" for f in beds),
       "and harvesting prose in the raised bed is not, so the check stays usable")


def check_callcard(yard):
    """Traps found in the record, one per genus, never one half of a pair."""
    from lib import bundle, callcard

    data = bundle.build(SCREENS)
    task = next(t for t in data["tasks"] if t["id"] == "t101")
    traps, dropped = callcard.find_traps(SCREENS, data, task)
    by_genus = {t["genus"]: t for t in traps}

    ok(set(by_genus) == {"Carex", "Yucca", "Asclepias"},
       "one trap per genus the record actually pairs up", sorted(by_genus))
    ok(by_genus["Carex"]["buy"] == ["Carex planostachys"]
       and by_genus["Carex"]["refuse"] == ["Carex texensis"],
       "the wanted species and the refused one land on the right sides")
    ok(by_genus["Asclepias"]["buy"]
       == ["Asclepias asperula", "Asclepias tuberosa"],
       "a record accepting two species keeps both as things to buy",
       by_genus["Asclepias"]["buy"])
    ok(by_genus["Yucca"]["refuse"] == ["Yucca recurvifolia"],
       "and an abbreviated refusal is expanded", by_genus["Yucca"]["refuse"])

    # "Same trap as the Yucca above" contributes a Yucca and nothing to buy.
    ok(all(t["buy"] and t["refuse"] for t in traps),
       "no panel is published with only one half of a trap")
    ok([n for n, _ in dropped] == ["Autumn sage"],
       "a plant named in the call and dropped from the design is reported once",
       dropped)

    page = callcard.build(SCREENS, data=data, link_images=True)
    ok(page and "Carex texensis" in page and "REFUSE THIS" in page,
       "and the page says so in as many words")


def check_pages(yard):
    """Every screen builds, and they point at each other rather than repeating.

    Built through `site.build_all` rather than one renderer at a time, because
    the property under test is that they form a set: the bar only offers a
    screen whose file is on disk, and the index only advertises what exists.
    Rendering each in isolation would pass while the set stayed broken.
    """
    from lib import site

    made, _ = site.build_all(SCREENS, link_images=True)
    ok({"INDEX.html", "TASKS.html", "CALENDAR.html", "WEEK.html",
        "PLAN.html", "CALL-CARD.html"} <= set(made),
       "one command builds every screen", sorted(made))

    def page(name):
        with open(os.path.join(yard, name), encoding="utf-8") as fh:
            return fh.read()

    tasks_page = page("TASKS.html")
    ok('id="t102"' in tasks_page,
       "every job has its own address on the task page")
    ok('href="PLAN.html#1-sowing"' in tasks_page,
       "and its method is a link into the document, not a copy of it")
    ok("How to sow." not in tasks_page,
       "the prose that explains a job stays in the one file that holds it")

    ok('href="TASKS.html#t102"' in page("CALENDAR.html"),
       "the calendar points at the job rather than restating it")

    index = page("INDEX.html")
    ok('href="TASKS.html"' in index and 'nav class="screens"' in index,
       "the index reaches the set, and carries the same bar as every page")
    ok('nav class="screens"' in page("PLAN.html"),
       "and so does a published markdown document")

    # This fixture has nothing dated to the current week, which is the point:
    # an empty week still has to belong to the set rather than dead-end.
    wk = page("WEEK.html")
    ok('href="TASKS.html"' in wk and 'nav class="screens"' in wk,
       "a week with nothing in it still reaches the rest of the set")
    ok('data-task="t102"' in tasks_page and 'data-task="t101"' in tasks_page,
       "a task carries its id as the checkbox key, so a tick means one thing")

    # A link to a doubt is the one link nothing else proves, because the card
    # it points at is written by a different module onto a different page.
    ok('href="INDEX.html#d01"' in tasks_page,
       "a job held up by a question links to the card")
    ok('id="d01"' in index,
       "and the index gives that card the anchor the link asks for")
    ok('href="INDEX.html#d02"' not in tasks_page
       and "settled question d02" in tasks_page,
       "a card already settled is named rather than linked, because it has "
       "left the index")


def check_standing_ticks(yard):
    """One standing job is one tick, on whichever day it is read.

    A repeat draws a box on every day it asks for work, so the same job holds
    several boxes on one page. Each carries the same id, which is what makes
    the tick mean one thing when it comes home — and what makes ticking one
    box leave its twins looking undone unless the script keeps them together.
    """
    from lib import week as W

    monday = datetime.date(2026, 10, 19)
    monday -= datetime.timedelta(days=monday.weekday())
    page = W.render_week_html(SCREENS, monday, today=monday)
    if not isinstance(page, str) or "<" not in page:
        with open(os.path.join(yard, "WEEK.html"), encoding="utf-8") as fh:
            page = fh.read()

    boxes = page.count('data-task="t104"')
    ok(boxes > 1, f"a daily job draws a box on each day it asks for work "
                  f"(drew {boxes})")
    ok("function twins(" in page and "twins(id).forEach" in page,
       "and ticking one of them ticks the rest, rather than only the one "
       "that was clicked")
    ok("function each_job(" in page and "each_job(function" in page,
       "the tally counts the job once, not once per day it lands on")


def main():
    global verbose
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-v", "--verbose", action="store_true")
    verbose = ap.parse_args().verbose

    root = tempfile.mkdtemp(prefix="yard-week-test-")
    os.environ["GARDEN_ROOT"] = root
    for mod in [m for m in list(sys.modules) if m.startswith("lib")]:
        del sys.modules[mod]
    from lib import week

    print("lib.week — does the calendar refuse over a plan it disagrees with\n")
    try:
        yard = make_yard(root)
        print(" a clean record")
        check_clean(week, yard)
        print("\n the digest")
        check_unstamped(week, yard)
        check_uncited(week, yard)
        check_source_is_hashed_whatever_the_file(week, yard)
        check_reference_is_not_hashed(week, yard)
        check_dotted_anchor(week, yard)
        check_digest(week, yard)
        print("\n the dates")
        check_date(week, yard)
        print("\n --force")
        check_force(week, yard)
        print("\n the way out")
        check_cheap_paths(week, yard)
        print("\n the ticks")
        check_sync(week, yard)
        print("\n the published table widths")
        check_column_widths()

        screens = make_screens(root)
        print("\n one anchor scheme")
        check_one_slug(screens)
        print("\n which plant a line names")
        check_matcher(screens)
        print("\n the bundle every page renders from")
        check_bundle(screens)
        print("\n what --links reports")
        check_link_report(screens)
        print("\n the traps the record states")
        check_callcard(screens)
        print("\n the set of screens")
        check_pages(screens)
        print("\n one standing job, one tick")
        check_standing_ticks(screens)
    finally:
        shutil.rmtree(root, ignore_errors=True)

    bad = [r for r in results if r[0] == "FAIL"]
    print(f"\n{len(results) - len(bad)} passed, {len(bad)} failed")
    if bad:
        print("\nThe calendar is not holding:")
        for _, label, _ in bad:
            print(f"  - {label}")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
